import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
import os

# Helvetica не содержит кириллицу: без TTF русские буквы рисуются квадратами.
FONT_NAME = 'Helvetica'
REGULAR_FONTS = [
    'C:/Windows/Fonts/arial.ttf',
    '/System/Library/Fonts/Supplemental/Arial.ttf',
    '/Library/Fonts/Arial Unicode.ttf',
    '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
]
BOLD_FONTS = [
    'C:/Windows/Fonts/arialbd.ttf',
    '/System/Library/Fonts/Supplemental/Arial Bold.ttf',
    '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
]

for fp in REGULAR_FONTS:
    if os.path.exists(fp):
        try:
            pdfmetrics.registerFont(TTFont('CyrillicFont', fp))
            FONT_NAME = 'CyrillicFont'
            break
        except Exception:
            pass

bold_registered = False
for fp in BOLD_FONTS:
    if os.path.exists(fp):
        try:
            pdfmetrics.registerFont(TTFont('CyrillicFontBold', fp))
            bold_registered = True
            break
        except Exception:
            pass

if FONT_NAME == 'CyrillicFont':
    bold_face = 'CyrillicFontBold' if bold_registered else 'CyrillicFont'
    pdfmetrics.registerFontFamily(
        'CyrillicFont',
        normal='CyrillicFont',
        bold=bold_face,
        italic='CyrillicFont',
        boldItalic=bold_face,
    )

METRIC_TITLES = {
    'speed_to_lead': 'Скорость первого ответа',
    'stagnation': 'Зависшие сделки',
    'discount_leakage': 'Утечка скидок',
    'sales_cycle': 'Цикл сделки',
    'key_account_risk': 'Зависимость от крупных клиентов',
    'dormant': 'Забытые клиенты',
    'funnel_dropoff': 'Провал воронки',
}

VERDICT_RU = {
    'critical': 'Критично',
    'watch': 'Следить',
    'ok': 'Норма',
    'skipped': 'Не посчитано',
}

STATUS_RU = {
    'new': 'Новая',
    'in_progress': 'В работе',
    'proposal': 'КП отправлено',
    'negotiation': 'Переговоры',
    'won': 'Выиграна',
    'lost': 'Проиграна',
    'other': 'Другое',
}


def money_ru(value) -> str:
    try:
        amount = float(value or 0)
    except (TypeError, ValueError):
        return '0 руб.'
    text = f'{amount:,.2f}'.replace(',', ' ')
    return f'{text} руб.'


def metric_ru(metric_id: str) -> str:
    return METRIC_TITLES.get(metric_id, metric_id)


def verdict_ru(verdict: str) -> str:
    key = (verdict or 'ok').lower()
    return VERDICT_RU.get(key, verdict or 'Норма')


def status_ru(status: str) -> str:
    key = (status or '').lower()
    return STATUS_RU.get(key, status or '—')


def generate_excel_report(snapshot, deals_qs) -> io.BytesIO:
    wb = Workbook()
    
    # Sheet 1: Резюме аудита
    ws_summary = wb.active
    ws_summary.title = "Резюме аудита"
    ws_summary.views.sheetView[0].showGridLines = True

    # Styling colors
    NAVY = "1E293B"
    BLUE = "2563EB"
    RED = "DC2626"
    GRAY_BG = "F8FAFC"
    WHITE = "FFFFFF"

    # Header title
    ws_summary.merge_cells("A1:F1")
    title_cell = ws_summary["A1"]
    title_cell.value = "X-RAY: АУДИТ ВОРОНКИ ПРОДАЖ"
    title_cell.font = Font(size=16, bold=True, color=WHITE)
    title_cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 40

    # Key Metrics Blocks
    totals = snapshot.totals or {}
    total_deals = totals.get("deals", 0)
    total_amount = totals.get("amount", "0")
    findings = snapshot.findings or []

    # Calculate money at risk
    money_at_risk = 0.0
    for f in findings:
        if f.get("money_impact"):
            try:
                money_at_risk += float(f["money_impact"])
            except:
                pass

    # Score calculation (100 - penalties)
    score = 100
    for f in findings:
        v = f.get("verdict")
        if v == "critical":
            score -= 30
        elif v == "watch":
            score -= 15
    score = max(score, 10)

    stats = [
        ("Снимок ID", str(snapshot.id)[:18] + "..."),
        ("Файл выгрузки", snapshot.filename or "ecommerce.csv"),
        ("Балл здоровья воронки", f"{score} / 100"),
        ("Всего сделок в базе", f"{total_deals} шт."),
        ("Совокупный объем сделок", money_ru(total_amount)),
        ("Денег под угрозой", money_ru(money_at_risk)),
    ]

    ws_summary.cell(row=3, column=1, value="КЛЮЧЕВЫЕ ПОКАЗАТЕЛИ").font = Font(bold=True, size=12, color=BLUE)

    r = 4
    for label, val in stats:
        c1 = ws_summary.cell(row=r, column=1, value=label)
        c2 = ws_summary.cell(row=r, column=2, value=val)
        c1.font = Font(bold=True, color="334155")
        c2.font = Font(bold=(label == "Денег под угрозой"), color=(RED if label == "Денег под угрозой" else "000000"))
        ws_summary.row_dimensions[r].height = 22
        r += 1

    # Headline
    r += 1
    ws_summary.cell(row=r, column=1, value="ГЛАВНЫЙ ВЫВОД АНАЛИТИКА:").font = Font(bold=True, size=12, color=BLUE)
    r += 1
    ws_summary.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=6)
    headline = str(snapshot.headline or "По воронке нет критических замечаний.").replace(" ₽", " руб.").replace("₽", " руб.")
    hl_cell = ws_summary.cell(row=r, column=1, value=headline)
    hl_cell.font = Font(size=12, italic=True, bold=True)
    hl_cell.fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    hl_cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws_summary.row_dimensions[r].height = 24
    ws_summary.row_dimensions[r+1].height = 24

    # Top threats table
    r += 3
    ws_summary.cell(row=r, column=1, value="НАЙДЕННЫЕ УГРОЗЫ И ПЛАН ДЕЙСТВИЙ").font = Font(bold=True, size=12, color=BLUE)
    r += 1

    headers = ["Метрика", "Уровень риска", "Потери в деньгах", "Порог проблемы", "Что сделать прямо сейчас"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=r, column=col_idx, value=h)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=BLUE, end_color=BLUE, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[r].height = 26

    for f in findings:
        r += 1
        v = verdict_ru(f.get("verdict") or "ok")
        impact = money_ru(f.get("money_impact") or 0)
        action = str(f.get("action") or "").replace(" ₽", " руб.").replace("₽", " руб.")
        ws_summary.cell(row=r, column=1, value=metric_ru(f.get("metric_id", "")))
        c_risk = ws_summary.cell(row=r, column=2, value=v)
        c_risk.font = Font(bold=True, color=(RED if v == "Критично" else "D97706"))
        ws_summary.cell(row=r, column=3, value=impact)
        ws_summary.cell(row=r, column=4, value=f.get("threshold_label", ""))
        action_c = ws_summary.cell(row=r, column=5, value=action)
        action_c.alignment = Alignment(wrap_text=True)
        ws_summary.row_dimensions[r].height = 36

    ws_summary.column_dimensions["A"].width = 20
    ws_summary.column_dimensions["B"].width = 16
    ws_summary.column_dimensions["C"].width = 22
    ws_summary.column_dimensions["D"].width = 30
    ws_summary.column_dimensions["E"].width = 60
    ws_summary.column_dimensions["F"].width = 15

    # Sheet 2: Реестр зависших сделок
    ws_deals = wb.create_sheet(title="Зависшие сделки")
    ws_deals.views.sheetView[0].showGridLines = True

    deal_headers = ["ID сделки", "Клиент / Контрагент", "Сумма сделки (руб.)", "Статус", "Менеджер", "Телефон / Контакт", "Дата создания"]
    for col_idx, h in enumerate(deal_headers, 1):
        c = ws_deals.cell(row=1, column=col_idx, value=h)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws_deals.row_dimensions[1].height = 28

    stagnant_deals = deals_qs.filter(status__in=["new", "in_progress", "proposal", "negotiation", "other"]).order_by("-amount")[:250]

    for row_idx, d in enumerate(stagnant_deals, 2):
        ws_deals.cell(row=row_idx, column=1, value=d.deal_id)
        ws_deals.cell(row=row_idx, column=2, value=d.client or "—")
        c_amt = ws_deals.cell(row=row_idx, column=3, value=float(d.amount))
        c_amt.number_format = "#,##0.00"
        ws_deals.cell(row=row_idx, column=4, value=status_ru(d.status))
        ws_deals.cell(row=row_idx, column=5, value=d.manager or "Не назначен")
        ws_deals.cell(row=row_idx, column=6, value=d.contact or "—")
        ws_deals.cell(row=row_idx, column=7, value=d.created_at.strftime("%d.%m.%Y") if d.created_at else "—")
        ws_deals.row_dimensions[row_idx].height = 20

    ws_deals.column_dimensions["A"].width = 18
    ws_deals.column_dimensions["B"].width = 30
    ws_deals.column_dimensions["C"].width = 20
    ws_deals.column_dimensions["D"].width = 16
    ws_deals.column_dimensions["E"].width = 24
    ws_deals.column_dimensions["F"].width = 22
    ws_deals.column_dimensions["G"].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf_report(snapshot, deals_qs) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        fontName=FONT_NAME,
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=12
    )

    h2_style = ParagraphStyle(
        'DocH2',
        fontName=FONT_NAME,
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#2563EB'),
        spaceBefore=14,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'DocBody',
        fontName=FONT_NAME,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155')
    )

    bold_body = ParagraphStyle(
        'DocBodyBold',
        fontName=FONT_NAME,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#0F172A')
    )

    callout_style = ParagraphStyle(
        'DocCallout',
        fontName=FONT_NAME,
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1E3A8A')
    )

    story = []

    # Header
    story.append(Paragraph("<b>X-RAY: АУДИТ ВОРОНКИ ПРОДАЖ</b>", title_style))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#2563EB'), spaceAfter=14))

    # Health Score & Money at Risk calculation
    findings = snapshot.findings or []
    money_at_risk = 0.0
    for f in findings:
        if f.get("money_impact"):
            try:
                money_at_risk += float(f["money_impact"])
            except:
                pass

    score = 100
    for f in findings:
        v = f.get("verdict")
        if v == "critical":
            score -= 30
        elif v == "watch":
            score -= 15
    score = max(score, 10)

    totals = snapshot.totals or {}
    total_deals = totals.get("deals", 0)
    total_amount = float(totals.get("amount", 0) or 0)

    # Summary Info Table
    summary_data = [
        [
            Paragraph(f"<b>Балл здоровья бизнеса:</b> {score} / 100", bold_body),
            Paragraph(f"<b>Сделок в отчете:</b> {total_deals} шт.", bold_body),
        ],
        [
            Paragraph(f"<b>Выручка в воронке:</b> {money_ru(total_amount)}", bold_body),
            Paragraph(f"<b>Денег под угрозой:</b> <font color='#DC2626'>{money_ru(money_at_risk)}</font>", bold_body)
        ]
    ]

    t_summary = Table(summary_data, colWidths=[260, 260])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_summary)
    story.append(Spacer(1, 14))

    # Headline
    story.append(Paragraph("<b>Главный диагноз:</b>", h2_style))
    hl_text = str(snapshot.headline or "Критических утечек в воронке не зафиксировано.").replace(" ₽", " руб.").replace("₽", " руб.")
    t_hl = Table([[Paragraph(hl_text, callout_style)]], colWidths=[520])
    t_hl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EFF6FF')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#BFDBFE')),
        ('PADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(t_hl)
    story.append(Spacer(1, 14))

    # Findings Table
    story.append(Paragraph("<b>Найденные угрозы и план действий:</b>", h2_style))
    
    findings_table_data = [
        [
            Paragraph("<b>Метрика</b>", bold_body),
            Paragraph("<b>Уровень</b>", bold_body),
            Paragraph("<b>Потери</b>", bold_body),
            Paragraph("<b>Что сделать прямо сейчас</b>", bold_body),
        ]
    ]

    for f in findings:
        v = verdict_ru(f.get("verdict") or "ok")
        v_color = "#DC2626" if v == "Критично" else "#D97706"
        impact = money_ru(f.get("money_impact") or 0)
        action = str(f.get("action") or "").replace(" ₽", " руб.").replace("₽", " руб.")
        findings_table_data.append([
            Paragraph(metric_ru(f.get("metric_id", "")), body_style),
            Paragraph(f"<font color='{v_color}'><b>{v}</b></font>", body_style),
            Paragraph(impact, bold_body),
            Paragraph(action, body_style),
        ])

    t_findings = Table(findings_table_data, colWidths=[100, 75, 95, 250])
    t_findings.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_findings)
    story.append(Spacer(1, 16))

    # Top Stagnant Deals Sample
    story.append(Paragraph("<b>Топ-5 самых крупных зависших сделок:</b>", h2_style))
    stagnant_sample = deals_qs.filter(status__in=["new", "in_progress", "proposal", "negotiation", "other"]).order_by("-amount")[:5]

    deals_table_data = [
        [
            Paragraph("<b>Клиент</b>", bold_body),
            Paragraph("<b>Сумма</b>", bold_body),
            Paragraph("<b>Менеджер</b>", bold_body),
            Paragraph("<b>Контакт</b>", bold_body),
        ]
    ]

    for d in stagnant_sample:
        deals_table_data.append([
            Paragraph(d.client or "—", body_style),
            Paragraph(money_ru(d.amount), bold_body),
            Paragraph(d.manager or "—", body_style),
            Paragraph(d.contact or "—", body_style),
        ])

    t_deals = Table(deals_table_data, colWidths=[170, 100, 130, 120])
    t_deals.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('PADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_deals)

    doc.build(story)
    buffer.seek(0)
    return buffer
